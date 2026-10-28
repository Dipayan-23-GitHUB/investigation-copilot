from __future__ import annotations
import csv
import hashlib
import io
import json
from datetime import date, datetime
from typing import Any
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from ..catalog import get_module
from ..models import Ingestion, SourceSheet, SourceRow

class IngestionError(ValueError):
    pass

def _json_value(value: Any):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)

def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())

def _map_sheet(module_cfg: dict, sheet_name: str):
    normalized = sheet_name.strip().upper()
    for item in module_cfg.get("sheets", []):
        expected = item["name"].strip().upper()
        if normalized == expected or normalized.startswith(expected + " ") or normalized.startswith(expected + "("):
            return item
    return None

def _records_from_xlsx(content: bytes, module_cfg: dict, header_row: int):
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise IngestionError(f"Could not read Excel workbook: {exc.__class__.__name__}") from exc
    output = []
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        headers = None
        for idx, row in enumerate(rows_iter, start=1):
            if idx == header_row:
                headers = [str(x).strip() if x is not None else "" for x in row]
                break
        if not headers:
            output.append({"sheet_name": ws.title, "headers": [], "rows": [], "source_object": "Unmapped source", "warnings": [f"Header row {header_row} was not found."]})
            continue
        last_nonempty = max((i for i, val in enumerate(headers) if not _empty(val)), default=-1)
        headers = headers[:last_nonempty + 1]
        if not headers or all(not x for x in headers):
            output.append({"sheet_name": ws.title, "headers": [], "rows": [], "source_object": "Unmapped source", "warnings": ["No column headers found on configured header row."]})
            continue
        clean_headers = []
        used = {}
        for i, h in enumerate(headers):
            name = h or f"column_{i+1}"
            # Keep headers stable, but disambiguate duplicates rather than silently overwriting values.
            used[name] = used.get(name, 0) + 1
            clean_headers.append(name if used[name] == 1 else f"{name}__{used[name]}")
        records = []
        for row_num, row in enumerate(rows_iter, start=header_row + 1):
            row_values = list(row[:len(clean_headers)])
            if not row_values or all(_empty(v) for v in row_values):
                continue
            if len(row_values) < len(clean_headers):
                row_values += [None] * (len(clean_headers) - len(row_values))
            records.append({clean_headers[i]: _json_value(row_values[i]) for i in range(len(clean_headers))})
        mapped = _map_sheet(module_cfg, ws.title)
        warnings = []
        if mapped:
            missing = [col for col in mapped.get("columns", []) if col not in clean_headers]
            if missing:
                warnings.append("Missing expected columns: " + ", ".join(missing))
        else:
            warnings.append("Sheet name is not mapped to a named object in the Source Ledger catalog; data is stored without a verified object mapping.")
        output.append({"sheet_name": ws.title, "headers": clean_headers, "rows": records, "source_object": mapped["object"] if mapped else "Unmapped source", "warnings": warnings})
    wb.close()
    return output

def _records_from_csv(content: bytes, header_row: int):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionError("CSV must be UTF-8 encoded.") from exc
    lines = text.splitlines()
    if len(lines) < header_row:
        raise IngestionError(f"Header row {header_row} is beyond the end of the CSV.")
    reader = csv.DictReader(lines[header_row - 1:])
    records = [{str(k): _json_value(v) for k, v in row.items()} for row in reader if any(not _empty(v) for v in row.values())]
    headers = reader.fieldnames or []
    return [{"sheet_name": "CSV", "headers": headers, "rows": records, "source_object": "Unmapped source", "warnings": ["CSV source mapping should be confirmed before using it for automated investigations."]}]

def _records_from_json(content: bytes):
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except Exception as exc:
        raise IngestionError("JSON file is not valid UTF-8 JSON.") from exc
    blocks = []
    if isinstance(payload, list):
        payload = {"records": payload}
    if not isinstance(payload, dict):
        raise IngestionError("JSON must be an array of records or an object containing arrays of records.")
    for name, value in payload.items():
        if not isinstance(value, list) or not all(isinstance(row, dict) for row in value):
            continue
        headers = list(dict.fromkeys(k for row in value for k in row.keys()))
        rows = [{str(k): _json_value(row.get(k)) for k in headers} for row in value]
        blocks.append({"sheet_name": str(name)[:160], "headers": headers, "rows": rows, "source_object": "Unmapped source", "warnings": ["JSON source mapping should be confirmed before using it for automated investigations."]})
    if not blocks:
        raise IngestionError("No array of record objects was found in the JSON payload.")
    return blocks

def ingest_file(db: Session, *, module_key: str, filename: str, content: bytes, header_row: int = 5, tenant_id: str = "demo") -> tuple[Ingestion, bool]:
    module_cfg = get_module(module_key)
    if not module_cfg:
        raise IngestionError(f"Unknown module_key '{module_key}'. Get valid values from GET /api/v1/catalog/modules.")
    if not content:
        raise IngestionError("Uploaded file is empty.")
    digest = hashlib.sha256(content).hexdigest()
    existing = db.query(Ingestion).filter_by(tenant_id=tenant_id, module_key=module_key, sha256=digest).first()
    if existing:
        return existing, True
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "xlsx":
        sheet_blocks = _records_from_xlsx(content, module_cfg, header_row)
    elif suffix == "csv":
        sheet_blocks = _records_from_csv(content, header_row)
    elif suffix == "json":
        sheet_blocks = _records_from_json(content)
    else:
        raise IngestionError("Unsupported file type. Use .xlsx, .csv, or .json; legacy .xls is not supported by this first ingestion API.")
    ingest_warnings = []
    expected_names = [s["name"].upper() for s in module_cfg.get("sheets", [])]
    found_names = [b["sheet_name"].strip().upper() for b in sheet_blocks]
    for expected in expected_names:
        if not any(n == expected or n.startswith(expected + " ") or n.startswith(expected + "(") for n in found_names):
            ingest_warnings.append(f"Expected sheet/object '{expected}' was not found in the upload.")
    for block in sheet_blocks:
        ingest_warnings.extend([f"{block['sheet_name']}: {w}" for w in block.get("warnings", [])])
    ingestion = Ingestion(id=str(uuid4()), tenant_id=tenant_id, module_key=module_key, filename=filename[:255], sha256=digest, status="completed_with_warnings" if ingest_warnings else "completed", row_count=sum(len(b["rows"]) for b in sheet_blocks), sheet_count=len(sheet_blocks), warnings=ingest_warnings)
    db.add(ingestion)
    for block in sheet_blocks:
        sheet = SourceSheet(id=str(uuid4()), ingestion_id=ingestion.id, sheet_name=block["sheet_name"], source_object=block["source_object"], headers=block["headers"], row_count=len(block["rows"]))
        db.add(sheet)
        db.flush()
        db.add_all([SourceRow(id=str(uuid4()), sheet_id=sheet.id, row_number=i, data=row) for i, row in enumerate(block["rows"], start=1)])
    db.commit()
    db.refresh(ingestion)
    return ingestion, False
